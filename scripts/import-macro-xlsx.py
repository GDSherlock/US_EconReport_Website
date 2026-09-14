"""One-time, read-only workbook adapter. Produces a reviewable update batch."""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import openpyxl

def extract(workbook, mapping, edition, file_hash):
    sheet = workbook[mapping['sheet']]
    if sheet[mapping['headerCell']].value != mapping['expectedHeader']:
        raise ValueError(f"Header mismatch: {mapping['seriesId']}")
    points, issues = {}, []
    date_col = openpyxl.utils.column_index_from_string(mapping['dateColumn'])
    value_col = openpyxl.utils.column_index_from_string(mapping['valueColumn'])
    for row in sheet.iter_rows(min_row=mapping['startRow'], max_col=max(date_col, value_col)):
        date, value = row[date_col-1].value, row[value_col-1].value
        if not isinstance(date, dt.datetime):
            if value is not None:
                issues.append({'cell': getattr(row[value_col-1], 'coordinate', ''), 'reason': 'non-date row'})
            continue
        cell = f"{mapping['valueColumn']}{row[date_col-1].row}"
        if date.date().isoformat() > edition:
            issues.append({'cell': cell, 'reason': 'period after source edition'})
            continue
        if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)):
            issues.append({'cell': cell, 'reason': f'invalid cached value: {value}'})
            value = None
        frequency = mapping['frequency']
        period = date.strftime('%Y-%m') if frequency == 'monthly' else f'{date.year}-Q{(date.month-1)//3+1}' if frequency == 'quarterly' else date.date().isoformat()
        point = {'period': period, 'value': value, 'status': 'missing' if value is None else 'observed', 'sourceRef': 'workbook', 'sourcePublishedAt': None}
        if period in points:
            if points[period]['value'] != value:
                raise ValueError(f"Conflicting period: {mapping['seriesId']} {period}")
            issues.append({'cell': cell, 'reason': 'exact duplicate'})
        points[period] = point
    data = sorted(points.values(), key=lambda p: p['period'])
    # Retain internal gaps; do not publish leading/trailing padding as observations.
    while data and data[0]['value'] is None: data.pop(0)
    while data and data[-1]['value'] is None: data.pop()
    if not data: raise ValueError(f"No data: {mapping['seriesId']}")
    series = {key: mapping[key] for key in ['seriesId', 'name', 'frequency', 'unit', 'adjustment', 'measure']}
    series.update(sources={'workbook': {'label': 'Wind、Bloomberg', 'locator': f"{mapping['sheet']}!{mapping['dateColumn']}{mapping['startRow']}:{mapping['valueColumn']}{sheet.max_row}；表头 {mapping['headerCell']}", 'retrievedAt': dt.date.today().isoformat(), 'fileSha256': file_hash, 'url': None}}, points=data)
    samples = [data[0], data[len(data)//2], data[-1]]
    return series, {'seriesId': mapping['seriesId'], 'count': len(data), 'missing': sum(p['value'] is None for p in data), 'samples': samples, 'issues': issues}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True)
    parser.add_argument('--edition', required=True)
    parser.add_argument('--out', required=True)
    parser.add_argument('--map', default=str(Path(__file__).with_name('macro-bootstrap-map.json')))
    parser.add_argument('--release-id', default='bootstrap-20260911')
    args = parser.parse_args()
    dt.date.fromisoformat(args.edition)
    config = json.loads(Path(args.map).read_text())
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    digest = hashlib.sha256(Path(args.input).read_bytes()).hexdigest()
    extracted = [extract(workbook, m, args.edition, digest) for m in config['series']]
    batch = {'releaseId': args.release_id, 'sourceEdition': args.edition, 'newSeries': [s for s, _ in extracted], 'chartDefinitions': config['charts'], 'updates': []}
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(batch, ensure_ascii=False), encoding='utf8')
    out.with_suffix('.audit.json').write_text(json.dumps([a for _, a in extracted], ensure_ascii=False, indent=2), encoding='utf8')
    print(f'Extracted {len(extracted)} series, {len(config["charts"])} charts; batch {out}')

if __name__ == '__main__': main()
